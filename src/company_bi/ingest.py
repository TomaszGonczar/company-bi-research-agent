"""Small CSV/XLSX readers and independent deterministic identity-row outcomes."""

import csv
from collections.abc import Sequence
from datetime import UTC, datetime
from pathlib import Path
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile
from zoneinfo import ZoneInfo

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from company_bi.models import BatchResult, CompanyIdentity, InputRow, Source
from company_bi.nip import InvalidNIP, validate_nip
from company_bi.registry import RegistryLookupError, lookup_company


class InputFileError(ValueError):
    """The file cannot safely supply a mandatory NIP column."""


def _nip_column(headers: Sequence[object]) -> int:
    columns = [
        index
        for index, value in enumerate(headers)
        if isinstance(value, str) and value.strip().casefold() == "nip"
    ]
    if len(columns) != 1:
        raise InputFileError("Input must have exactly one nip header (case-insensitive)")
    return columns[0]


def read_input(path: Path) -> list[tuple[int, object]]:
    """Row numbers include the header; CSV numbers count logical records."""
    try:
        if path.suffix.lower() == ".csv":
            with path.open(encoding="utf-8-sig", newline="") as file:
                reader = csv.reader(file, strict=True)
                column = _nip_column(next(reader, []))
                return [
                    (number, row[column] if column < len(row) else None)
                    for number, row in enumerate(reader, start=2)
                ]
        if path.suffix.lower() == ".xlsx":
            workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
            try:
                sheet = workbook.worksheets[0]
                rows = sheet.iter_rows(values_only=True)
                column = _nip_column(next(rows, ()))
                return [
                    (number, row[column] if column < len(row) else None)
                    for number, row in enumerate(rows, start=2)
                ]
            finally:
                workbook.close()
        raise InputFileError("Only UTF-8 CSV and XLSX input files are supported")
    except (
        OSError,
        UnicodeError,
        csv.Error,
        BadZipFile,
        InvalidFileException,
        ParseError,
    ) as error:
        raise InputFileError(f"Cannot read input file: {error}") from error


def resolve_file(path: Path) -> list[BatchResult]:
    rows = read_input(path)
    as_of = datetime.now(ZoneInfo("Europe/Warsaw")).date()
    cache: dict[str, tuple[CompanyIdentity | None, Source] | RegistryLookupError] = {}
    results: list[BatchResult] = []
    for number, value in rows:
        input_nip = "" if value is None else str(value)
        try:
            row = InputRow(row_number=number, nip=validate_nip(value))
        except InvalidNIP as error:
            results.append(
                BatchResult(
                    row_number=number,
                    input_nip=input_nip,
                    status="invalid_input",
                    error_code="INVALID_NIP",
                    reason=str(error),
                    completed_at=datetime.now(UTC),
                )
            )
            continue
        if row.nip not in cache:
            try:
                cache[row.nip] = lookup_company(row.nip, as_of)
            except RegistryLookupError as error:
                cache[row.nip] = error
        outcome = cache[row.nip]
        if isinstance(outcome, RegistryLookupError):
            results.append(
                BatchResult(
                    row_number=number,
                    input_nip=input_nip,
                    nip=row.nip,
                    status="failed",
                    error_code=outcome.code,
                    reason=str(outcome),
                    completed_at=datetime.now(UTC),
                )
            )
            continue
        identity, source = outcome
        if identity is None:
            results.append(
                BatchResult(
                    row_number=number,
                    input_nip=input_nip,
                    nip=row.nip,
                    status="unresolved",
                    sources=[source],
                    error_code="COMPANY_NOT_FOUND",
                    reason="No entity returned by the MF VAT register for this NIP and query date",
                    completed_at=datetime.now(UTC),
                )
            )
        else:
            results.append(
                BatchResult(
                    row_number=number,
                    input_nip=input_nip,
                    nip=row.nip,
                    status="resolved",
                    identity=identity,
                    sources=[source],
                    completed_at=datetime.now(UTC),
                )
            )
    return results
